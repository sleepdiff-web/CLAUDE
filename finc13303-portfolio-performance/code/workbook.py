"""
Build the supporting Excel workbook for the NDIF six-year performance review.

Every table behind the presentation, plus the raw monthly return series, so a
reader can verify any figure in the deck without running the Python.
"""
import json, os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)          # repository sub-directory root
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "deliverables")
os.makedirs(OUT, exist_ok=True)
R = json.load(open(os.path.join(DATA, "results.json")))
ST, SS, REG = R["stats"], R["stock_stats"], R["reg"]

SER = pd.read_csv(os.path.join(DATA, "series_monthly.csv"), index_col=0)
STK = pd.read_csv(os.path.join(DATA, "stocks_monthly.csv"), index_col=0)
ALLR = pd.read_csv(os.path.join(DATA, "all_monthly_returns.csv"), index_col=0)
WP = pd.read_csv(os.path.join(DATA, "weight_path.csv"), index_col=0)
ANN = pd.read_csv(os.path.join(DATA, "annual.csv"), index_col=0)

NAVY = "FF0B2545"; GOLD = "FFC9A227"; TINT = "FFF2F6FC"; WHITE = "FFFFFFFF"
HEAD = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor=TINT)
HFONT = Font(name="Calibri", size=10.5, bold=True, color=WHITE)
BFONT = Font(name="Calibri", size=10.5)
LFONT = Font(name="Calibri", size=10.5, bold=True)
TFONT = Font(name="Calibri", size=13, bold=True, color="FF0B2545")
NOTE = Font(name="Calibri", size=9, italic=True, color="FF52514E")
THIN = Side(style="thin", color="FFE6E5E1")
BOX = Border(bottom=THIN)

wb = Workbook()
wb.remove(wb.active)

N, Q, S = "NDIF", "NASDAQ-100 (QQQ)", "S&P 500 (SPY)"
EW, MV = "Equal weight (1/N)", "MV optimised (ex-ante 2014-16)"
SIXTY = "60/40 (SPY/AGG)"
TICK = list(STK.columns)
GICS = {"MSFT": "Information Technology", "GOOGL": "Communication Services",
        "AMZN": "Consumer Discretionary", "V": "Financials", "EQIX": "Real Estate",
        "ILMN": "Health Care", "ROK": "Industrials", "NEE": "Utilities",
        "COST": "Consumer Staples", "ALB": "Materials", "SLB": "Energy"}
COMPANY = {"MSFT": "Microsoft", "GOOGL": "Alphabet", "AMZN": "Amazon", "V": "Visa",
           "EQIX": "Equinix", "ILMN": "Illumina", "ROK": "Rockwell Automation",
           "NEE": "NextEra Energy", "COST": "Costco", "ALB": "Albemarle",
           "SLB": "Schlumberger (SLB)"}
TIER = {"MSFT": 1, "GOOGL": 1, "AMZN": 1, "V": 2, "EQIX": 2, "ILMN": 2, "ROK": 2,
        "NEE": 3, "COST": 3, "ALB": 3, "SLB": 3}
SECTOR_ETF = {"MSFT": "XLK", "GOOGL": "XLC", "AMZN": "XLY", "V": "XLF", "EQIX": "XLRE",
              "ILMN": "XLV", "ROK": "XLI", "NEE": "XLU", "COST": "XLP", "ALB": "XLB",
              "SLB": "XLE"}

def sheet(name, title, subtitle=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = TFONT
    if subtitle:
        ws["A2"] = subtitle; ws["A2"].font = NOTE
    ws.freeze_panes = "A5"
    return ws

def write_table(ws, top, rows, widths=None, pcts=None, nums=None, header=True,
                first_bold=True):
    """rows[0] is the header. pcts/nums are column indices (0-based)."""
    pcts, nums = pcts or [], nums or []
    for i, row in enumerate(rows):
        r = top + i
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.font = HFONT if (i == 0 and header) else BFONT
            c.alignment = Alignment(horizontal="left" if j == 0 else "right",
                                    vertical="center", wrap_text=(i == 0))
            if i == 0 and header:
                c.fill = HEAD
            else:
                if i % 2 == 1:
                    c.fill = BAND
                c.border = BOX
                if first_bold and j == 0:
                    c.font = LFONT
                if j in pcts and isinstance(v, (int, float)):
                    c.number_format = "0.0%"
                elif j in nums and isinstance(v, (int, float)):
                    c.number_format = "0.00"
        ws.row_dimensions[r].height = 30 if (i == 0 and header) else 15
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w
    return top + len(rows)

# ------------------------------------------------------------------ 1. Read me
ws = sheet("Read me", "Northpoint Digital Innovation Fund — six-year performance workbook",
           "Supporting calculations for the Assignment 2 portfolio performance presentation")
info = [["Item", "Detail"],
        ["Student", "Callum O'Connor"],
        ["Student ID", "14053836"],
        ["Subject", "FINC13-303 Portfolio Analysis and Investments"],
        ["Assessment", "Assignment 2, Part 1 — Portfolio performance online presentation"],
        ["Period analysed", "January 2017 – December 2022 (72 monthly observations)"],
        ["Mandate", "US$100,000,000"],
        ["Holdings", "11 US-listed equities, one per GICS sector"],
        ["Rebalancing", "Quarterly to target weights (24 rebalances)"],
        ["Primary benchmark", "NASDAQ-100 total return (QQQ ETF)"],
        ["Secondary benchmark", "S&P 500 total return (SPY ETF)"],
        ["Risk-free rate", "Fama–French one-month Treasury bill series"],
        ["Risk factors", "Kenneth R. French Data Library — Mkt-RF, SMB, HML, RMW, CMA, MOM"],
        ["Regression method", "OLS with Newey–West (HAC) standard errors, 4 lags"],
        ["Price basis", "Dividend- and split-adjusted monthly closing prices"],
        ["Excluded", "Transaction costs, market impact, taxes, cash drag"],
        ["", ""],
        ["Sheet", "Contents"],
        ["Holdings", "Target weights, tiers, and per-holding six-year statistics"],
        ["Performance summary", "Full metric table for the fund, the variants and the benchmarks"],
        ["Annual returns", "Calendar-year returns for every series"],
        ["Regressions", "CAPM, FF3, FF5 and FF5+momentum output"],
        ["Attribution", "Brinson–Fachler decomposition and per-holding contribution"],
        ["Seasonality", "Monthly return grid by year and calendar-month averages"],
        ["Monthly returns", "The 72-month return series for every strategy and benchmark"],
        ["Stock returns", "The 72-month return series for the eleven holdings"],
        ["Weight path", "Portfolio weights at the start of each month, as run"],
        ["Factor data", "Fama–French factors and the risk-free rate, as used"]]
write_table(ws, 4, info, widths=[26, 92])

# ------------------------------------------------------------------ 2. Holdings
ws = sheet("Holdings", "Holdings, target weights and six-year statistics",
           "Statistics computed on monthly total returns, January 2017 – December 2022")
rows = [["Ticker", "Company", "GICS sector", "Tier", "Target weight", "Total return",
         "CAGR", "Volatility", "Sharpe", "Sortino", "Max drawdown", "Beta",
         "CAPM alpha p.a.", "Alpha t-stat", "Sector ETF", "Sector ETF return",
         "Excess vs sector", "Contribution to fund"]]
for t in TICK:
    rows.append([t, COMPANY[t], GICS[t], f"Tier {TIER[t]}", R["weights_target"][t],
                 R["stock_total"][t], SS[t]["cagr"], SS[t]["vol"], SS[t]["sharpe"],
                 SS[t]["sortino"], SS[t]["max_dd"], SS[t]["beta"],
                 R["reg_stock_capm"][t]["alpha_a"], R["reg_stock_capm"][t]["alpha_t"],
                 SECTOR_ETF[t], R["etf_total"][t],
                 R["stock_total"][t] - R["etf_total"][t], R["contrib"][t]])
rows.append(["NDIF", "Portfolio", "All eleven", "—", 1.0, ST[N]["cum_return"],
             ST[N]["cagr"], ST[N]["vol"], ST[N]["sharpe"], ST[N]["sortino"],
             ST[N]["max_dd"], ST[N]["beta"], REG["CAPM"]["alpha_a"],
             REG["CAPM"]["alpha_t"], "—", R["attrib"]["sector_bench_cum"],
             ST[N]["cum_return"] - R["attrib"]["sector_bench_cum"], ST[N]["cum_return"]])
write_table(ws, 4, rows, widths=[9, 20, 22, 7, 12, 12, 9, 11, 8, 8, 13, 7, 14, 11, 11, 15, 14, 16],
            pcts=[4, 5, 6, 7, 10, 12, 15, 16, 17], nums=[8, 9, 11, 13])
ws.cell(row=4 + len(rows) + 1, column=1,
        value="Contribution is the compounded contribution of each holding to the fund's "
              "179.2% cumulative return, using the realised weight path.").font = NOTE

# ----------------------------------------------------- 3. Performance summary
ws = sheet("Performance summary", "Full performance and risk metrics",
           "72 monthly observations, January 2017 – December 2022")
keys = [(N, "NDIF (as run)"), ("NDIF (buy & hold)", "NDIF buy & hold"),
        (EW, "Equal weight 1/N"), (MV, "Mean-variance (2014-16)"),
        ("MV optimised (perfect foresight)", "MV perfect foresight"),
        (Q, "NASDAQ-100 (QQQ)"), (S, "S&P 500 (SPY)"), (SIXTY, "60/40 SPY/AGG")]
metrics = [("Observations (months)", "months", "int"),
           ("Cumulative return", "cum_return", "pct"),
           ("Annualised return (CAGR)", "cagr", "pct"),
           ("Annualised volatility", "vol", "pct"),
           ("Mean monthly return", "mean_m", "pct"),
           ("Median monthly return", "median_m", "pct"),
           ("Best month", "best", "pct"), ("Worst month", "worst", "pct"),
           ("Skewness", "skew", "num"), ("Excess kurtosis", "kurt", "num"),
           ("Jarque-Bera p-value", "jb_p", "num"),
           ("Positive months", "n_pos", "int"), ("Negative months", "n_neg", "int"),
           ("Hit rate", "hit", "pct"), ("Average positive month", "avg_pos", "pct"),
           ("Average negative month", "avg_neg", "pct"),
           ("Downside deviation", "downside_dev", "pct"),
           ("Maximum drawdown", "max_dd", "pct"),
           ("Sharpe ratio", "sharpe", "num"), ("Sortino ratio", "sortino", "num"),
           ("Treynor ratio", "treynor", "num"), ("Calmar ratio", "calmar", "num"),
           ("Information ratio (vs S&P 500)", "info_ratio", "num"),
           ("M-squared (risk-matched return)", "m2", "pct"),
           ("Beta (vs S&P 500)", "beta", "num"),
           ("Correlation (vs S&P 500)", "corr", "num"),
           ("Tracking error (vs S&P 500)", "te", "pct"),
           ("Upside capture (vs S&P 500)", "up_capture", "pct"),
           ("Downside capture (vs S&P 500)", "down_capture", "pct"),
           ("VaR 95% historical (monthly)", "var95_hist", "pct"),
           ("VaR 95% parametric normal", "var95_param", "pct"),
           ("VaR 95% Cornish-Fisher", "var95_cf", "pct"),
           ("CVaR 95% (expected shortfall)", "cvar95_hist", "pct"),
           ("VaR 99% historical (monthly)", "var99_hist", "pct"),
           ("CVaR 99% (expected shortfall)", "cvar99_hist", "pct"),
           ("Terminal value on $100m (US$)", "terminal", "usd")]
rows = [["Metric"] + [lbl for _, lbl in keys]]
pcts, nums = [], []
for i, (label, key, kind) in enumerate(metrics):
    rows.append([label] + [ST[k].get(key) for k, _ in keys])
write_table(ws, 4, rows, widths=[32] + [16] * len(keys))
for i, (label, key, kind) in enumerate(metrics):
    r = 5 + i
    for j in range(len(keys)):
        c = ws.cell(row=r, column=j + 2)
        c.number_format = {"pct": "0.00%", "num": "0.000", "int": "0",
                           "usd": "$#,##0"}[kind]

# ------------------------------------------------------------- 4. Annual returns
ws = sheet("Annual returns", "Calendar-year total returns",
           "Geometrically compounded from monthly returns")
rows = [["Series"] + [str(int(y)) for y in ANN.index] + ["Cumulative", "CAGR"]]
for k, lbl in keys:
    rows.append([lbl] + [float(ANN[k][y]) for y in ANN.index] +
                [ST[k]["cum_return"], ST[k]["cagr"]])
write_table(ws, 4, rows, widths=[26] + [12] * (len(ANN.index) + 2),
            pcts=list(range(1, len(ANN.index) + 3)))

# --------------------------------------------------------------- 5. Regressions
ws = sheet("Regressions", "Factor-model output — Jensen's alpha",
           "Dependent variable: fund return in excess of the one-month T-bill. "
           "Newey–West (HAC) standard errors, 4 lags.")
facs = ["Mkt-RF", "SMB", "HML", "RMW", "CMA", "MOM"]
rows = [["Model", "Alpha (monthly)", "Alpha (annualised)", "Alpha t-stat", "Alpha p-value"]
        + sum([[f, f + " t-stat"] for f in facs], []) + ["R-squared", "Adjusted R-squared",
                                                         "Observations"]]
for m in ["CAPM", "FF3", "FF5", "FF6"]:
    g = REG[m]
    row = [{"CAPM": "CAPM (single factor)", "FF3": "Fama-French 3-factor",
            "FF5": "Fama-French 5-factor", "FF6": "FF5 + momentum (Carhart)"}[m],
           g["alpha_m"], g["alpha_a"], g["alpha_t"], g["alpha_p"]]
    for f in facs:
        row += [g.get(f), g.get(f + "_t")]
    row += [g["r2"], g["r2adj"], g["n"]]
    rows.append(row)
g = R["reg_qqq"]["FF5"]
row = ["NASDAQ-100 benchmark (FF5)", g["alpha_m"], g["alpha_a"], g["alpha_t"], g["alpha_p"]]
for f in facs:
    row += [g.get(f), g.get(f + "_t")]
rows.append(row + [g["r2"], g["r2adj"], g["n"]])
write_table(ws, 4, rows, widths=[26, 14, 16, 12, 12] + [11] * 12 + [11, 16, 12],
            pcts=[1, 2], nums=list(range(3, 20)))
top = 4 + len(rows) + 2
ws.cell(row=top, column=1, value="Per-holding CAPM regressions").font = TFONT
rows2 = [["Ticker", "Alpha (annualised)", "Alpha t-stat", "Beta", "Beta t-stat", "R-squared"]]
for t in TICK:
    g = R["reg_stock_capm"][t]
    rows2.append([t, g["alpha_a"], g["alpha_t"], g["Mkt-RF"], g["Mkt-RF_t"], g["r2"]])
write_table(ws, top + 2, rows2, pcts=[1], nums=[2, 3, 4, 5])

# --------------------------------------------------------------- 6. Attribution
ws = sheet("Attribution", "Performance attribution",
           "Brinson–Fachler decomposition against an equal-weight benchmark of the eleven "
           "GICS sector ETFs, summed arithmetically over 72 months")
a = R["attrib"]
rows = [["Component", "Contribution (points of return)", "Share of active return"],
        ["Sector allocation", a["allocation"], a["allocation"] / a["total_active_arith"]],
        ["Security selection", a["selection"], a["selection"] / a["total_active_arith"]],
        ["Interaction", a["interaction"], a["interaction"] / a["total_active_arith"]],
        ["Total active return", a["total_active_arith"], 1.0]]
end = write_table(ws, 4, rows, widths=[30, 28, 22], pcts=[1, 2])
rows2 = [["Reference", "Cumulative return"],
         ["NDIF portfolio", ST[N]["cum_return"]],
         ["Equal-weight sector-ETF benchmark", a["sector_bench_cum"]],
         ["NASDAQ-100", ST[Q]["cum_return"]],
         ["S&P 500", ST[S]["cum_return"]]]
end = write_table(ws, end + 2, rows2, pcts=[1])
rows3 = [["Ticker", "Company", "Contribution (points)", "Share of fund return",
          "Stock total return", "Sector ETF total return"]]
for t in sorted(TICK, key=lambda x: -R["contrib"][x]):
    rows3.append([t, COMPANY[t], R["contrib"][t], R["contrib_pct"][t],
                  R["stock_total"][t], R["etf_total"][t]])
write_table(ws, end + 2, rows3, pcts=[2, 3, 4, 5])

# --------------------------------------------------------------- 7. Seasonality
ws = sheet("Seasonality", "Seasonality of fund returns",
           "NDIF monthly return by calendar month and year, and the calendar-month average")
grid = pd.DataFrame(R["monthly_grid"]).astype(float)
grid.index = grid.index.astype(int); grid = grid.sort_index()
grid.columns = [int(c) for c in grid.columns]; grid = grid[sorted(grid.columns)]
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
rows = [["Year"] + MON + ["Full year"]]
for y in grid.index:
    v = list(grid.loc[y].values)
    rows.append([int(y)] + v + [float(np.prod([1 + x for x in v]) - 1)])
ma = pd.DataFrame(R["month_avg"])
ma.index = ma.index.astype(int); ma = ma.sort_index()
rows.append(["Average — NDIF"] + list(ma[N].values) + [""])
rows.append(["Average — NASDAQ-100"] + list(ma[Q].values) + [""])
rows.append(["Average — S&P 500"] + list(ma[S].values) + [""])
write_table(ws, 4, rows, widths=[20] + [9] * 13, pcts=list(range(1, 14)))

# ------------------------------------------------------- 8-11. raw return series
def dump(name, df, title, sub, fmt="0.00%"):
    ws = sheet(name, title, sub)
    rows = [["Month"] + list(df.columns)]
    for idx, r in df.iterrows():
        rows.append([str(idx)] + [None if pd.isna(x) else float(x) for x in r.values])
    write_table(ws, 4, rows, widths=[11] + [14] * len(df.columns), first_bold=False)
    for i in range(len(df)):
        for j in range(len(df.columns)):
            ws.cell(row=5 + i, column=j + 2).number_format = fmt

dump("Monthly returns", SER, "Monthly total returns — strategies and benchmarks",
     "January 2017 – December 2022")
dump("Stock returns", STK, "Monthly total returns — the eleven holdings",
     "Dividend- and split-adjusted, January 2017 – December 2022")
dump("Weight path", WP, "Portfolio weights at the start of each month",
     "Quarterly rebalancing to target; weights drift within each quarter")
dump("Factor data", ALLR[["Mkt-RF", "SMB", "HML", "RMW", "CMA", "MOM", "RF"]],
     "Fama–French risk factors and the risk-free rate",
     "Kenneth R. French Data Library, monthly US research series", fmt="0.0000")

path = os.path.join(OUT, "FINC13303_Ass2_PerformanceWorkbook_14053836_OConnor.xlsx")
wb.save(path)
print("saved", path)
print("sheets:", wb.sheetnames)
